from io import BytesIO

from openpyxl import load_workbook

from app.db import connect
from app.services.excel_io import export_workbook, preview_students, student_template
from app.services.reports import ReportService
from app.services.students import StudentService


def test_manual_attendance_and_alpha_summary(tmp_path):
    db_path = tmp_path / "report.db"
    students = StudentService(db_path)
    reports = ReportService(db_path)
    first = students.save_student(None, "Hadir", 1, {})
    second = students.save_student(None, "Alpha", 2, {})
    reports.save_manual(first, "2026-10-04", "Hadir", "07:00", "Tepat waktu")
    summary = reports.summary("2026-10-04", "2026-10-04")
    assert summary == [
        {"nomor_urut": 1, "nama": "Hadir", "Hadir": 1, "Izin": 0, "Sakit": 0, "Alpha": 0},
        {"nomor_urut": 2, "nama": "Alpha", "Hadir": 0, "Izin": 0, "Sakit": 0, "Alpha": 1},
    ]
    reports.save_manual(second, "2026-10-04", "Izin", None, "Sakit ringan")
    assert reports.summary("2026-10-04", "2026-10-04")[1]["Izin"] == 1
    with connect(db_path) as connection:
        attendance_id = connection.execute(
            "SELECT id FROM absensi WHERE santri_id = ?", (second,)
        ).fetchone()[0]
    reports.delete_manual(attendance_id)
    assert reports.summary("2026-10-04", "2026-10-04")[1]["Alpha"] == 1
    alpha_rows = reports.attendance_rows("2026-10-04", "2026-10-04", status="Alpha")
    assert [(row["nama"], row["status"]) for row in alpha_rows] == [("Alpha", "Alpha")]


def test_export_round_trip_has_two_sheets_and_custom_field(tmp_path):
    content = export_workbook(
        [
            {
                "tanggal": "2026-10-04",
                "nomor_urut": 1,
                "nama": "Alya",
                "waktu": "07:00",
                "status": "Hadir",
                "sumber": "manual",
                "catatan": "Baik",
                "extra_data": {"nama_wali": "Ibu Alya"},
            }
        ],
        [{"nomor_urut": 1, "nama": "Alya", "Hadir": 1, "Izin": 0, "Sakit": 0, "Alpha": 0}],
        [{"field_key": "nama_wali", "field_label": "Nama Wali"}],
    )
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames == ["Rekap", "Ringkasan"]
    assert workbook["Rekap"].freeze_panes == "A2"
    assert workbook["Rekap"]["L1"].value == "Nama Wali"
    assert workbook["Rekap"]["L2"].value == "Ibu Alya"
    assert workbook["Ringkasan"]["D2"].value == 0


def test_import_preview_valid_and_problem_rows(tmp_path):
    students = StudentService(tmp_path / "import.db")
    students.add_field("Nama Wali")
    workbook = load_workbook(BytesIO(student_template(students.fields(True))))
    sheet = workbook.active
    sheet.append([1, "Valid", "Wali Valid"])
    sheet.append([2, "", "Tanpa Nama"])
    sheet.append([3, "Valid", "Duplikat"])
    output = BytesIO()
    workbook.save(output)

    preview = preview_students(output.getvalue(), students.fields(True))
    assert len(preview["rows"]) == 1
    assert len(preview["problems"]) == 2
    assert preview["rows"][0]["nama"] == "Valid"
