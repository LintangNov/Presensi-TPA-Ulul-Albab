from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _fit_columns(sheet) -> None:
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(width, 40)


def export_workbook(
    rows: list[dict[str, Any]],
    summary: list[dict[str, Any]],
    fields: list[Any],
) -> bytes:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "Rekap"
    field_headers = [field["field_label"] for field in fields]
    headers = [
        "Tanggal",
        "Nomor",
        "Nama",
        "Paket",
        "Jam Masuk",
        "Jam Pulang",
        "Status Pulang",
        "Denda (Menit)",
        "Status",
        "Sumber",
        "Catatan",
    ] + field_headers
    detail.append(headers)
    for cell in detail[1]:
        cell.font = Font(bold=True)
    for row in rows:
        extra = row.get("extra_data") or {}
        detail.append(
            [
                row.get("tanggal", ""),
                row.get("nomor_urut", ""),
                row.get("nama", ""),
                row.get("paket", "A"),
                row.get("waktu", "") or "-",
                row.get("waktu_pulang", "") or "-",
                row.get("status_pulang", "") or "-",
                row.get("denda_menit", 0),
                row.get("status", ""),
                row.get("sumber", ""),
                row.get("catatan", "") or "",
            ]
            + [extra.get(field["field_key"], "") for field in fields]
        )
    detail.freeze_panes = "A2"
    _fit_columns(detail)

    sheet = workbook.create_sheet("Ringkasan")
    sheet.append(["Nomor", "Nama", "Hadir", "Izin", "Sakit", "Alpha"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in summary:
        sheet.append(
            [row["nomor_urut"], row["nama"], row["Hadir"], row["Izin"], row["Sakit"], row["Alpha"]]
        )
    sheet.freeze_panes = "A2"
    _fit_columns(sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def student_template(fields: list[Any]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Murid"
    headers = ["Nomor", "Nama", "Paket"] + [field["field_label"] for field in fields]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    _fit_columns(sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def preview_students(content: bytes, fields: list[Any]) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    headers = [str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(values_only=True))]
    required = {"Nomor", "Nama"}
    if not required.issubset(headers):
        raise ValueError("Header Excel harus memiliki kolom Nomor dan Nama.")
    indexes = {header: index for index, header in enumerate(headers)}
    rows: list[dict[str, Any]] = []
    problems: list[str] = []
    seen_names: set[str] = set()
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(value not in (None, "") for value in values):
            continue
        name = str(values[indexes["Nama"]]).strip() if values[indexes["Nama"]] is not None else ""
        number = values[indexes["Nomor"]]
        paket_val = "A"
        if "Paket" in indexes and values[indexes["Paket"]]:
            raw_p = str(values[indexes["Paket"]]).strip().upper()
            if raw_p in {"A", "B"}:
                paket_val = raw_p

        if not name:
            problems.append(f"Baris {row_number}: nama wajib diisi.")
            continue
        key = name.casefold()
        if key in seen_names:
            problems.append(f"Baris {row_number}: nama duplikat di file ({name}).")
            continue
        seen_names.add(key)
        extra = {
            field["field_key"]: str(values[indexes[field["field_label"]]] or "")
            for field in fields
            if field["field_label"] in indexes
        }
        rows.append({
            "nama": name,
            "nomor_urut": int(number) if number not in (None, "") else None,
            "paket": paket_val,
            "extra_data": extra,
        })
    return {"rows": rows, "problems": problems}
